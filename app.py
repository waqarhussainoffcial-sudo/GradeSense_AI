import random
import io
from datetime import datetime

import streamlit as st
import streamlit.components.v1 as components
from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.lib.units import inch
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER
from reportlab.platypus import (
    SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, HRFlowable
)
import tempfile

# Excel export (Phase 2.7 / STEP 1) — openpyxl is used only for building the
# .xlsx workbook. Import is wrapped so the app keeps working (minus the
# Excel button) even in an environment where openpyxl isn't installed yet.
try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False

# MongoDB integration (Phase 3.1) — wrapped so the rest of the app keeps
# working (minus Save/View Records) even before pymongo is installed or
# MONGO_URI is configured. See config.py / database.py / student_operations.py.
try:
    import config
    from database import test_connection
    from student_operations import (
        save_student_result,
        get_all_students,
        get_student_by_roll,
    )
    MONGODB_AVAILABLE = True
except ImportError:
    MONGODB_AVAILABLE = False

# Authentication (Phase 3.2 / STEP 2) — uses the functions already built
# in auth.py (verify_user). Wrapped the same way as the imports above so
# a missing auth.py fails gracefully with a message instead of crashing.
try:
    from auth import verify_user, get_user_by_username
    AUTH_AVAILABLE = True
except ImportError:
    AUTH_AVAILABLE = False

# GradeSense AI Insights Engine — AI feedback, trend predictor & chatbot.
# Wrapped the same way as the imports above so the app keeps working
# (minus the AI Insights section) even if this module is missing.
try:
    from ai_insights import (
        generate_ai_feedback,
        predict_next_performance,
        build_chat_context,
        answer_student_question,
    )
    AI_INSIGHTS_AVAILABLE = True
except ImportError:
    AI_INSIGHTS_AVAILABLE = False

# Dashboards (Phase 3.3 / STEP 5) — role-specific dashboards built in
# dashboard/admin_dashboard.py, dashboard/teacher_dashboard.py, and
# dashboard/student_dashboard.py (Steps 1-4). Imported here so app.py can
# route to the right dashboard based on the authenticated user's role.
from dashboard.admin_dashboard import render_admin_dashboard
from dashboard.teacher_dashboard import render_teacher_dashboard
from dashboard.student_dashboard import render_student_dashboard

# ==========================================================
# APPLICATION METADATA (Phase 2.7)
# Central place for app name/version so every panel (sidebar, About,
# footer, PDF/Excel exports) stays in sync when bumped in the future.
# ==========================================================
APP_NAME = "GradeSense AI"
APP_VERSION = "3.0"
APP_YEAR = datetime.now().year

# ==========================================================
# PAGE CONFIG (UI ONLY — does not affect any logic below)
# ==========================================================
st.set_page_config(
    page_title="GradeSense AI | Smart SIS Dashboard",
    page_icon="🎓",
    layout="wide",
    initial_sidebar_state="collapsed",
)

# --- Session State Initialization ---
if "result" not in st.session_state:
    st.session_state.result = False

if "name" not in st.session_state:
    st.session_state.name = ""

if "roll_no" not in st.session_state:
    st.session_state.roll_no = ""

if "institution_name" not in st.session_state:
    st.session_state.institution_name = ""

if "total" not in st.session_state:
    st.session_state.total = 0

if "percentage" not in st.session_state:
    st.session_state.percentage = 0.0

if "subjects" not in st.session_state:
    st.session_state.subjects = []

if "marks" not in st.session_state:
    st.session_state.marks = []

# ---------- Phase 2.7 / STEP 4: Settings-panel state ----------
# These only control *display* — they never touch calculation logic.
if "dev_mode" not in st.session_state:
    st.session_state.dev_mode = False

if "show_analytics" not in st.session_state:
    st.session_state.show_analytics = True

if "show_charts" not in st.session_state:
    st.session_state.show_charts = True

# ---------- Phase 3.2 / STEP 2: Authentication state ----------
if "logged_in" not in st.session_state:
    st.session_state.logged_in = False

if "username" not in st.session_state:
    st.session_state.username = ""

if "full_name" not in st.session_state:
    st.session_state.full_name = ""

if "role" not in st.session_state:
    st.session_state.role = ""


# ---------- Phase 3.2 / STEP 3: Role helpers (RBAC) ----------
# Single source of truth for role checks so access rules aren't
# duplicated/hardcoded in multiple places across the file.
def is_admin():
    return st.session_state.role == "Admin"


def is_teacher():
    return st.session_state.role == "Teacher"


def is_student():
    return st.session_state.role == "Student"


def has_calculator_access():
    """Admin and Teacher can access the full calculator/reports/analytics/
    database features. Student cannot."""
    return st.session_state.role in ("Admin", "Teacher")


def render_login_screen():
    """Phase 3.2 / STEP 2 — Professional Login Screen.

    Renders a centered Username/Password/Login form and authenticates
    via auth.verify_user(). On success, stores logged_in/username/
    full_name/role in session_state and reruns. On failure, shows a
    Streamlit error. Purely an auth gate — does not touch any
    calculator/PDF/Excel/analytics logic.
    """
    st.markdown(
        """
        <div style="text-align:center; padding: 40px 0 10px 0;">
            <div style="font-size: 48px; line-height:1;"></div>
            <div style="font-weight:700; font-size: 1.4rem; margin-top:6px;">
                GradeSense AI
            </div>
            <div style="font-size:0.9rem; color:#888;">Please sign in to continue</div>
        </div>
        """,
        unsafe_allow_html=True,
    )

    _, login_col, _ = st.columns([1, 1.2, 1])
    with login_col:
        with st.container(border=True):
            st.subheader("Login")

            if not AUTH_AVAILABLE:
                st.error(
                    "Authentication module ('auth.py') isn't available in "
                    "this environment. Please contact the administrator."
                )
                return

            login_username = st.text_input("Username", key="login_username_input")
            login_password = st.text_input(
                "Password", type="password", key="login_password_input"
            )

            if st.button("Login", use_container_width=True, type="primary"):
                # ---------- Phase 3.2 / STEP 4-5 BUGFIX: MongoDB-backed login ----------
                # Reuses the same MONGODB_AVAILABLE / test_connection()
                # pattern already used elsewhere in this file (Save to
                # Database, Student Records) — no new connection is
                # created. verify_user() is the single, primary
                # credential check: it takes username + password
                # together, uses bcrypt internally, and never returns
                # or exposes password_hash.
                if not MONGODB_AVAILABLE:
                    st.warning(
                        "MongoDB integration isn't available in this "
                        "environment — install `pymongo` to enable login."
                    )
                else:
                    db_connected, db_conn_msg = test_connection()
                    if not db_connected:
                        st.warning(f"Could not reach the database: {db_conn_msg}")
                    else:
                        user = verify_user(login_username, login_password)

                        # ---------- TEMPORARY DEBUG — remove after diagnosis ----------
                        st.write("DEBUG — repr(login_username):", repr(login_username))
                        st.write("DEBUG — repr(login_password):", repr(login_password))
                        st.write(
                            "DEBUG — verify_user() returned:",
                            "None" if user is None else f"user document with keys: {list(user.keys())}"
                        )
                        print(f"[LOGIN DEBUG] login_username={login_username!r} "
                              f"verify_user_result={'None' if user is None else user}")
                        # ---------- END TEMPORARY DEBUG ----------

                        if user:
                            st.session_state.logged_in = True
                            st.session_state.username = user["username"]
                            st.session_state.full_name = user["full_name"]
                            st.session_state.role = user["role"]
                            st.rerun()
                        else:
                            # verify_user() returns None both for wrong
                            # credentials and for a disabled account, so
                            # do a secondary lookup only here (failure
                            # path) to show the disabled-account message
                            # specifically, without gating the primary
                            # credential check behind it.
                            existing_user = get_user_by_username(login_username)
                            if existing_user is not None and existing_user.get("status") != "active":
                                st.error("This account has been disabled.")
                            else:
                                st.error("Invalid username or password.")


def reset_all_data():
    """Phase 2.7 / STEP 3 — Reset All Data.

    Clears every key Streamlit is tracking (form inputs, calculated
    results, charts, settings toggles) and forces a clean rerun so the
    app comes back up exactly like a fresh session. This does not touch
    any calculation/GPA logic — it only clears state.
    """
    for key in list(st.session_state.keys()):
        del st.session_state[key]
    st.rerun()


# ---------- Phase 3.2 / STEP 2: Login gate ----------
# Everything below (sidebar, calculator, analytics, exports, records)
# requires an authenticated session. Logged-out users only ever see
# the login screen above.
if not st.session_state.logged_in:
    render_login_screen()
    st.stop()


# ==========================================================
# Phase 2.7 / STEP 9: Light responsive/visual polish.
# Pure CSS — does not alter any widget behaviour or logic.
# ==========================================================
st.markdown(
    """
    <style>
        /* Tighter, more consistent spacing between top-level blocks */
        .block-container { padding-top: 2rem; padding-bottom: 2rem; }

        /* Give bordered containers ("cards") a softer, more professional look */
        div[data-testid="stVerticalBlockBorderWrapper"] {
            border-radius: 12px !important;
        }

        /* Sidebar polish */
        section[data-testid="stSidebar"] {
            border-right: 1px solid rgba(120, 120, 120, 0.25);
        }

        /* Centered footer */
        .sgc-footer {
            text-align: center;
            padding-top: 1.2rem;
            margin-top: 1.5rem;
            border-top: 1px solid rgba(120, 120, 120, 0.25);
            color: #888888;
            font-size: 0.85rem;
            line-height: 1.6;
        }
        .sgc-footer b { color: #666666; }
    </style>
    """,
    unsafe_allow_html=True,
)

# ==========================================================
# Phase 2.7 / STEP 4 + STEP 8: SIDEBAR
# Branding, navigation info, quick tips, and the Settings panel.
# Purely presentational — reads/writes only the settings keys above.
# ==========================================================
with st.sidebar:
    # ---------- Branding / Logo placeholder ----------
    st.markdown(
        """
        <div style="text-align:center; padding: 4px 0 10px 0;">
            <div style="font-size: 42px; line-height:1;"></div>
            <div style="font-weight:700; font-size: 1.05rem; margin-top:4px;">
                GradeSense AI
            </div>
            <div style="font-size:0.75rem; color:#888;">Student Information System</div>
        </div>
        """,
        unsafe_allow_html=True,
    )
    st.divider()

    # ---------- Navigation Information ----------
    st.markdown("**Navigation**")
    st.caption(
        "1⃣ Fill Student Information → 2⃣ Enter Subjects → "
        "3⃣ Calculate Grade → 4⃣ Review Report → "
        "5⃣ Export PDF / Excel / Print"
    )

    # ---------- Quick Tips ----------
    with st.expander("Quick Tips", expanded=False):
        st.markdown(
            "- Marks must be between **0 and 100**.\n"
            "- Subject names must be **unique**.\n"
            "- Choose **School** or **University** mode before entering marks.\n"
            "- Use **Reset All Data** below to start a fresh entry.\n"
            "- Toggle charts off in Settings for a faster, lighter view."
        )

    st.divider()

    # ---------- Settings Panel ----------
    st.markdown("**Settings**")

    with st.expander("Theme Information", expanded=False):
        st.caption(
            "This app follows Streamlit's active theme (light/dark). "
            "Change it from Streamlit's own menu → Settings → Theme."
        )

    st.text_input("Application Version", value=APP_VERSION, disabled=True)

    st.session_state.dev_mode = st.toggle(
        "Developer Mode", value=st.session_state.dev_mode,
        help="Shows extra internal/debug details for troubleshooting."
    )
    st.session_state.show_analytics = st.toggle(
        "Show Analytics Dashboard", value=st.session_state.show_analytics
    )
    st.session_state.show_charts = st.toggle(
        "Show Charts (Bar/Pie)", value=st.session_state.show_charts
    )

    with st.expander("Future Settings", expanded=False):
        st.caption(
            "Reserved for upcoming preferences (e.g. custom PDF branding, "
            "multi-language support, custom grading scales)."
        )

    # ---------- Phase 3.1: Database status ----------
    st.markdown("**Database Status**")
    if not MONGODB_AVAILABLE:
        st.caption("pymongo not installed")
    else:
        _sidebar_connected, _sidebar_msg = test_connection()
        if _sidebar_connected:
            st.caption("MongoDB connected")
        else:
            st.caption("MongoDB not reachable")

    st.divider()

    # ---------- Phase 3.2 / STEP 2: Account / Logout ----------
    st.markdown("**Account**")
    st.caption(f"{st.session_state.full_name or st.session_state.username} · {st.session_state.role}")
    if st.button("Logout", use_container_width=True):
        st.session_state.logged_in = False
        st.session_state.username = ""
        st.session_state.full_name = ""
        st.session_state.role = ""
        st.rerun()

    # ---------- Phase 3.3: Admin page navigation ----------
    # Visible to Admin only — Teacher and Student never see this.
    # Lets the Admin switch between separate pages (User Management,
    # Student Info & Subjects, Calculate Result, AI Insights, Student
    # Records) instead of one long scrolling page.
    if is_admin():
        st.markdown("**Admin Pages**")
        _admin_pages = [
            "User Management",
            "Student Info & Subjects",
            "Calculate Result",
            "AI Insights",
            "Student Records",
        ]
        if "admin_page" not in st.session_state:
            st.session_state.admin_page = _admin_pages[0]
        st.session_state.admin_page = st.radio(
            "Admin navigation",
            _admin_pages,
            index=_admin_pages.index(st.session_state.admin_page),
            label_visibility="collapsed",
        )

    st.divider()

    # ---------- Reset All Data ----------
    if st.button("Reset All Data", use_container_width=True):
        reset_all_data()

    if st.session_state.dev_mode:
        st.caption(f"Session keys: {len(st.session_state.keys())}")


# ==========================================================
# HEADER / ABOUT / HELP — used on the Teacher single-page flow.
# ==========================================================
def render_header_about_help():
    # ==========================================================
    # HEADER
    # ==========================================================
    st.title("GradeSense AI")
    st.caption("AI-Powered Student Information System · Report Card, Transcript & Smart Insights")

    # ---------- Phase 2.7 / STEP 5: About ----------
    with st.expander("ℹAbout this Application", expanded=False):
        st.markdown(f"""
    - **Application Name:** {APP_NAME}
    - **Version:** {APP_VERSION}
    - **Purpose:** Generate professional School Marksheets and University
      Academic Transcripts, complete with grading, GPA, analytics,
      AI-driven performance insights, and exportable PDF/Excel reports.

    **Technologies Used**
    - Python
    - Streamlit
    - ReportLab (PDF generation)
    - Plotly (interactive charts)
    - OpenPyXL (Excel export)
    - GradeSense AI Insights Engine (feedback, predictor & chatbot)

    **Developer Credits:** _Add your name / team here._
    """)

    # ---------- Phase 2.7 / STEP 6: Help & User Guide ----------
    with st.expander("Help & User Guide", expanded=False):
        st.markdown("""
    **1. Entering Marks**
    Fill in the student's name, roll number, and institution name, then enter
    marks for each subject (0–100).

    **2. School Mode**
    Select **School** as the institution type, choose class/section, and the
    app produces a board-style **Marksheet** with grade, result, and an
    automatic teacher's remark.

    **3. University Mode**
    Select **University**, choose department/semester, and the app produces
    an official **Academic Transcript** with GPA and academic standing.

    **4. Generate PDF**
    Click **Calculate Grade** — a formatted PDF (Marksheet or Transcript)
    is generated automatically and appears under **Download Report**.

    **5. Export Excel**
    In the same **Download Report** section, click **Download Excel
    Report** to get an .xlsx copy with the same information.

    **6. Analytics**
    Scroll down to **Student Performance Analytics** for summary metrics,
    bar/pie charts, and (for University) GPA analysis. Toggle these on/off
    from the sidebar **Settings**.

    **7. Reset**
    Use **Reset All Data** in the sidebar to clear the form and start over.
    """)

    st.divider()


def validate_inputs(name, roll_no, institution_name, subjects, marks):
    """Professional input validation (Phase 2.4).

    Returns a list of human-readable error messages. An empty list means
    the inputs are valid and calculation may proceed. Does not touch any
    grade/GPA/percentage calculation logic.
    """
    errors = []

    if not name or not name.strip():
        errors.append("Enter Student Name.")

    if not roll_no or not roll_no.strip():
        errors.append("Enter Roll Number.")

    if not institution_name or not institution_name.strip():
        errors.append("Enter Institution Name.")

    cleaned_subjects = [s.strip() if s else "" for s in subjects]
    if any(s == "" for s in cleaned_subjects):
        errors.append("Enter Subject Name.")

    if any(m is None for m in marks):
        errors.append("Marks must be between 0 and 100.")
    elif any(m < 0 or m > 100 for m in marks):
        errors.append("Marks must be between 0 and 100.")

    non_empty_subjects = [s.lower() for s in cleaned_subjects if s]
    if len(non_empty_subjects) != len(set(non_empty_subjects)):
        errors.append("Duplicate subject names are not allowed.")

    return errors



# ==========================================================
# STUDENT INFO + SUBJECT DETAILS + CALCULATE (Admin: one page;
# Teacher: part of the single scrolling page)
# ==========================================================
def render_data_entry_section():
    # ==========================================================
    # STEP 2 — STUDENT INFORMATION CARD
    # ==========================================================
    with st.container(border=True):
        st.subheader("Student Information")

        col_a, col_b = st.columns(2)
        with col_a:
            name = st.text_input("Enter Student Name")
            institution_type = st.selectbox(
                "Select Institution Type",
                ["School", "University"]
            )
        with col_b:
            roll_no = st.text_input("Enter Roll Number")
            institution_name = st.text_input("Enter Institution Name")

        st.markdown("&nbsp;", unsafe_allow_html=False)

        if institution_type == "School":
            st.markdown("**School Details**")
            col_c, col_d = st.columns(2)
            with col_c:
                student_class = st.selectbox("Select Class", ["9th", "10th", "11th", "12th"])
            with col_d:
                section = st.selectbox("Select Section", ["A", "B", "C", "D"])
            academic_session = st.text_input("Academic Session", value="2025-2026")

        if institution_type == "University":
            st.markdown("**University Details**")
            col_c, col_d = st.columns(2)
            with col_c:
                department = st.selectbox(
                    "Select Department",
                    ["B.Ed", "BS Computer Science", "BS English", "BBA", "B.Com"]
                )
            with col_d:
                semester = st.selectbox(
                    "Select Semester",
                    ["1st", "2nd", "3rd", "4th", "5th", "6th", "7th", "8th"]
                )
            academic_session = st.text_input("Academic Session", value="2025-2026", key="uni_academic_session_input")

    st.write("")

    # ==========================================================
    # STEP 3 — SUBJECTS CARD
    # ==========================================================
    with st.container(border=True):
        st.subheader("Subject Details")

        subject_count = st.selectbox("Select Number of Subjects", [3, 4, 5, 6, 7])

        subjects = []
        marks = []

        with st.expander("Subject Entry", expanded=True):
            for i in range(subject_count):
                col_s, col_m = st.columns([2, 1])
                with col_s:
                    subject_name = st.text_input(f"Enter Subject {i+1} Name", key=f"subject_name_{i}")
                with col_m:
                    subject_mark = st.number_input(
                        f"Enter {subject_name if subject_name else 'Subject'} Marks",
                        min_value=0,
                        max_value=100,
                        key=f"subject_mark_{i}"
                    )
                subjects.append(subject_name)
                marks.append(subject_mark)

    st.write("")
    calculate = st.button("Calculate Grade", use_container_width=True, type="primary")

    # --- Validate, then save everything into session_state when button is clicked ---
    if calculate:
        validation_errors = validate_inputs(name, roll_no, institution_name, subjects, marks)

        if validation_errors:
            st.session_state.result = False
            for err in validation_errors:
                st.error(f"{err}")
        else:
            st.session_state.result = True

            total = sum(marks)
            percentage = total / (subject_count * 100) * 100

            st.session_state.total = total
            st.session_state.percentage = percentage
            st.session_state.name = name
            st.session_state.roll_no = roll_no
            st.session_state.institution_name = institution_name
            st.session_state.subjects = subjects
            st.session_state.marks = marks
            st.session_state.institution_type = institution_type

            if institution_type == "School":
                st.session_state.student_class = student_class
                st.session_state.section = section
                st.session_state.academic_session = academic_session
                st.session_state.issue_date = datetime.now().strftime("%B %d, %Y")
                st.session_state.report_number = (
                    f"RPT-{datetime.now().strftime('%Y%m%d')}-{random.randint(1000, 9999)}"
                )
            else:
                st.session_state.department = department
                st.session_state.semester = semester
                st.session_state.university_academic_session = academic_session
                st.session_state.uni_issue_date = datetime.now().strftime("%B %d, %Y")
                st.session_state.transcript_number = (
                    f"TRX-{datetime.now().strftime('%Y%m%d')}-{random.randint(1000, 9999)}"
                )

            st.success("Result calculated successfully.")



def build_subject_rows(institution_type, subjects, marks):
    rows = []
    all_pass = True
    for i in range(len(subjects)):
        m = marks[i]
        passing = 60 if institution_type == "University" else 33

        if m >= 80:
            grade, gp = "A", 4.0
        elif m >= 70:
            grade, gp = "B", 3.0
        elif m >= 60:
            grade, gp = "C", 2.0
        elif m >= 50:
            grade, gp = "D", 1.0
        else:
            grade, gp = "F", 0.0

        status = "Pass" if m >= passing else "Fail"
        if status == "Fail":
            all_pass = False

        # Placeholder transcript identifiers — used only by the University
        # transcript view/PDF. They do not affect School output or any
        # existing grade/GPA/percentage calculations.
        subj_label = subjects[i] if subjects[i] else "SUB"
        course_code = f"{subj_label[:3].upper()}{101 + i}"
        credit_hours = 3

        rows.append({
            "subject": subjects[i],
            "max": 100,
            "passing": passing,
            "obtained": m,
            "grade": grade,
            "gp": gp,
            "status": status,
            "course_code": course_code,
            "credit_hours": credit_hours,
        })
    return rows, all_pass


def get_overall_grade(percentage):
    """Overall marksheet grade band for the School report (display only)."""
    if percentage >= 80:
        return "A"
    elif percentage >= 70:
        return "B"
    elif percentage >= 60:
        return "C"
    elif percentage >= 50:
        return "D"
    else:
        return "F"


def get_teacher_remark(percentage):
    """Automatic teacher remark for the School report (display only)."""
    if percentage >= 95:
        return "Outstanding"
    elif percentage >= 90:
        return "Excellent"
    elif percentage >= 80:
        return "Very Good"
    elif percentage >= 70:
        return "Good"
    elif percentage >= 60:
        return "Average"
    else:
        return "Needs Improvement"


def get_academic_standing(gpa):
    """Academic standing band for the University transcript (display only)."""
    if gpa >= 4.00:
        return "Outstanding"
    elif gpa >= 3.75:
        return "Excellent"
    elif gpa >= 3.25:
        return "Very Good"
    elif gpa >= 2.75:
        return "Good Standing"
    elif gpa >= 2.00:
        return "Satisfactory"
    else:
        return "Academic Warning"


def get_performance_indicator(percentage):
    """Overall performance classification for the analytics dashboard
    (Phase 2.6, display only). Returns (icon, label).
    """
    if percentage >= 90:
        return "", "Excellent"
    elif percentage >= 80:
        return "", "Very Good"
    elif percentage >= 70:
        return "", "Good"
    elif percentage >= 60:
        return "", "Average"
    else:
        return "", "Needs Improvement"


def render_analytics_dashboard(inst_type, subj_rows, overall_status, show_charts=True):
    """Professional analytics dashboard (Phase 2.6).

    Purely a display layer built on top of values that are already
    calculated elsewhere (subj_rows, st.session_state.percentage, etc.).
    Does not perform or alter any grade/GPA/percentage calculation.

    Phase 2.7: `show_charts` lets the Settings panel hide the Plotly
    bar/pie charts for a lighter view — everything else is unaffected.
    """
    st.divider()
    st.subheader("Student Performance Analytics")

    subject_names = [r["subject"] for r in subj_rows]
    obtained = [r["obtained"] for r in subj_rows]
    percentage = st.session_state.percentage
    overall_grade = get_overall_grade(percentage)

    passed_count = sum(1 for r in subj_rows if r["status"] == "Pass")
    failed_count = sum(1 for r in subj_rows if r["status"] == "Fail")
    highest_marks = max(obtained)
    lowest_marks = min(obtained)
    average_marks = sum(obtained) / len(obtained)
    highest_subject = subject_names[obtained.index(highest_marks)]
    lowest_subject = subject_names[obtained.index(lowest_marks)]

    # ---------- STEP 2: Summary Cards ----------
    with st.container(border=True):
        st.markdown("**Overview**")
        m1, m2, m3 = st.columns(3)
        m1.metric("Overall Percentage", f"{percentage:.2f}%")
        m2.metric("Overall Grade", overall_grade)
        m3.metric("Overall Result", overall_status)

        m4, m5, m6 = st.columns(3)
        m4.metric("Subject Count", len(subj_rows))
        m5.metric("Passed Subjects", passed_count)
        m6.metric("Failed Subjects", failed_count)

        m7, m8, m9 = st.columns(3)
        m7.metric("Highest Marks", highest_marks)
        m8.metric("Lowest Marks", lowest_marks)
        m9.metric("Average Marks", f"{average_marks:.2f}")

        # ---------- STEP 2 (University addition) ----------
        if inst_type == "University":
            total_gp = sum(r["gp"] for r in subj_rows)
            gpa = total_gp / len(subj_rows)
            total_credit_hours = sum(r["credit_hours"] for r in subj_rows)

            u1, u2, u3 = st.columns(3)
            u1.metric("Semester GPA", f"{gpa:.2f}")
            u2.metric("Total Grade Points", f"{total_gp:.1f}")
            u3.metric("Total Credit Hours", total_credit_hours)

    st.write("")

    # ---------- STEP 3 & 4: Bar Chart + Pie Chart ----------
    # Phase 2.7: hidden entirely when the sidebar "Show Charts" toggle is
    # off. The chart-building code itself is untouched.
    if not show_charts:
        st.caption("Charts are hidden (enable in sidebar Settings → Show Charts).")
    else:
        with st.expander("Subject Performance Chart", expanded=True):
            try:
                import plotly.express as px

                bar_fig = px.bar(
                    x=subject_names, y=obtained,
                    labels={"x": "Subject", "y": "Obtained Marks"},
                    title="Subject-wise Obtained Marks",
                    color=obtained, color_continuous_scale="Blues", text=obtained,
                )
                bar_fig.update_traces(
                    textposition="outside",
                    hovertemplate="<b>%{x}</b><br>Obtained: %{y}<extra></extra>",
                )
                bar_fig.update_layout(
                    showlegend=False, coloraxis_showscale=False,
                    margin=dict(t=60, l=10, r=10, b=10),
                )
                st.plotly_chart(bar_fig, use_container_width=True)

                pie_fig = px.pie(
                    names=subject_names, values=obtained,
                    title="Subject-wise Contribution to Total Marks",
                    hole=0.35,
                )
                pie_fig.update_traces(
                    textinfo="percent+label",
                    hovertemplate="<b>%{label}</b><br>Obtained: %{value} (%{percent})<extra></extra>",
                )
                pie_fig.update_layout(margin=dict(t=60, l=10, r=10, b=10))
                st.plotly_chart(pie_fig, use_container_width=True)
            except ImportError:
                st.warning(
                    "Interactive charts need the 'plotly' package. "
                    "Install it with `pip install plotly` to see the bar and pie charts here."
                )

    # ---------- STEP 5: Performance Analysis ----------
    with st.expander("Performance Statistics", expanded=False):
        p1, p2 = st.columns(2)
        with p1:
            st.success(f"**Highest Scoring Subject:** {highest_subject} ({highest_marks} marks)")
            st.info(f"**Average Marks:** {average_marks:.2f}")
        with p2:
            st.error(f"**Lowest Scoring Subject:** {lowest_subject} ({lowest_marks} marks)")
            st.info(f"**Top Performer Subject:** {highest_subject}")
        st.warning(f"**Weakest Subject:** {lowest_subject}")

    # ---------- STEP 6: Performance Indicator ----------
    icon, label = get_performance_indicator(percentage)
    indicator_message = f"{icon} **Performance Indicator:** {label}"
    if label in ("Excellent", "Very Good"):
        st.success(indicator_message)
    elif label in ("Good", "Average"):
        st.info(indicator_message)
    else:
        st.warning(indicator_message)

    # ---------- STEP 7: University GPA Analytics ----------
    if inst_type == "University":
        with st.expander("GPA Analysis", expanded=False):
            total_gp = sum(r["gp"] for r in subj_rows)
            gpa = total_gp / len(subj_rows)
            total_credit_hours = sum(r["credit_hours"] for r in subj_rows)
            standing = get_academic_standing(gpa)

            g1, g2, g3 = st.columns(3)
            g1.metric("Semester GPA", f"{gpa:.2f}")
            g2.metric("Total Grade Points", f"{total_gp:.1f}")
            g3.metric("Credit Hours", total_credit_hours)

            st.info(f"**Academic Standing:** {standing}")
            st.write("**GPA Progress** (scale: 0.00 – 4.00)")
            st.progress(min(gpa / 4.0, 1.0))


def draw_logo_placeholder(label, accent_hex):
    """Reusable bordered logo placeholder for PDF headers (Phase 2.5).

    Returns a small Table flowable that can sit next to a title. Swap the
    single cell's content for a reportlab Image(...) flowable once a real
    logo file is available — no other header code needs to change.
    """
    placeholder_style = ParagraphStyle(
        "LogoPlaceholder", fontName="Helvetica-Bold", fontSize=7,
        alignment=TA_CENTER, textColor=colors.HexColor(accent_hex),
        leading=9,
    )
    box = Table([[Paragraph(label, placeholder_style)]], colWidths=[50], rowHeights=[50])
    box.setStyle(TableStyle([
        ("BOX", (0, 0), (-1, -1), 1, colors.HexColor(accent_hex)),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
    ]))
    return box


def _draw_pdf_footer(canvas, doc, accent_hex):
    """Shared professional footer drawer used by both the School and
    University PDFs (Phase 2.5). Only the accent color differs between
    them, so the layout stays identical and future footer changes only
    need to happen in one place.
    """
    now = datetime.now()
    canvas.saveState()

    canvas.setStrokeColor(colors.HexColor("#cccccc"))
    canvas.setLineWidth(0.6)
    canvas.line(0.7 * inch, 0.74 * inch, letter[0] - 0.7 * inch, 0.74 * inch)

    canvas.setFont("Helvetica-Bold", 7.5)
    canvas.setFillColor(colors.HexColor(accent_hex))
    canvas.drawString(0.7 * inch, 0.60 * inch, "Generated by GradeSense AI")

    canvas.setFont("Helvetica", 7)
    canvas.setFillColor(colors.HexColor("#888888"))
    canvas.drawString(0.7 * inch, 0.48 * inch, f"Python \u2022 Streamlit \u2022 ReportLab  |  Version {APP_VERSION}")
    canvas.drawString(
        0.7 * inch, 0.36 * inch,
        f"Generated on {now.strftime('%B %d, %Y')} at {now.strftime('%I:%M %p')}"
    )
    canvas.drawRightString(letter[0] - 0.7 * inch, 0.36 * inch, f"Page {doc.page}")
    canvas.restoreState()


def add_school_footer(canvas, doc):
    """Professional footer for the School marksheet PDF."""
    _draw_pdf_footer(canvas, doc, "#1a1a2e")


def add_university_footer(canvas, doc):
    """Professional footer for the University transcript PDF."""
    _draw_pdf_footer(canvas, doc, "#12263a")


def generate_excel_report(inst_type, subj_rows, overall_status):
    """Phase 2.7 / STEP 1 — Professional Excel (.xlsx) export.

    Mirrors the PDF's structure (Student Information, Subject Table,
    Summary, Remarks/Academic Standing) using openpyxl. This is purely a
    display/export layer built on values already calculated elsewhere
    (subj_rows, st.session_state.total/percentage, etc.) — it does not
    perform or alter any grade/GPA/percentage calculation.

    Returns an in-memory BytesIO buffer ready for st.download_button, or
    None if openpyxl isn't installed in this environment.
    """
    if not OPENPYXL_AVAILABLE:
        return None

    wb = Workbook()
    ws = wb.active
    ws.title = "School Marksheet" if inst_type == "School" else "Academic Transcript"

    # ---- Reusable styles ----
    accent_hex = "1A1A2E" if inst_type == "School" else "12263A"
    title_font = Font(bold=True, size=14, color="FFFFFF")
    title_fill = PatternFill("solid", fgColor=accent_hex)
    section_font = Font(bold=True, size=12, color=accent_hex)
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor=accent_hex)
    label_font = Font(bold=True)
    thin_border = Border(*(Side(style="thin", color="CCCCCC"),) * 4)
    center = Alignment(horizontal="center", vertical="center")

    max_cols = 6 if inst_type == "School" else 8
    row = 1

    # ---- Title Block ----
    title_text = (
        f"{st.session_state.institution_name or 'School Name'} — Student Marksheet"
        if inst_type == "School"
        else f"{st.session_state.institution_name or 'University Name'} — Official Academic Transcript"
    )
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=max_cols)
    title_cell = ws.cell(row=row, column=1, value=title_text)
    title_cell.font = title_font
    title_cell.fill = title_fill
    title_cell.alignment = center
    ws.row_dimensions[row].height = 24
    row += 2

    # ---- STUDENT INFORMATION ----
    ws.cell(row=row, column=1, value="Student Information").font = section_font
    row += 1

    if inst_type == "School":
        info_pairs = [
            ("Student Name", st.session_state.name),
            ("Roll Number", st.session_state.roll_no),
            ("Class", st.session_state.student_class),
            ("Section", st.session_state.section),
            ("Institution", st.session_state.institution_name),
            ("Academic Session", st.session_state.academic_session),
            ("Issue Date", st.session_state.issue_date),
            ("Report No", st.session_state.report_number),
        ]
    else:
        info_pairs = [
            ("Student Name", st.session_state.name),
            ("Roll Number", st.session_state.roll_no),
            ("Department", st.session_state.department),
            ("Semester", st.session_state.semester),
            ("Institution", st.session_state.institution_name),
            ("Academic Session", st.session_state.university_academic_session),
            ("Issue Date", st.session_state.uni_issue_date),
            ("Transcript No", st.session_state.transcript_number),
        ]

    for label, value in info_pairs:
        ws.cell(row=row, column=1, value=label).font = label_font
        ws.cell(row=row, column=2, value=value)
        row += 1
    row += 1

    # ---- SUBJECT / COURSE TABLE ----
    ws.cell(
        row=row, column=1,
        value="Subject-wise Marks" if inst_type == "School" else "Course-wise Performance"
    ).font = section_font
    row += 1

    if inst_type == "School":
        headers = ["Subject", "Maximum Marks", "Passing Marks", "Obtained Marks", "Grade", "Result"]
    else:
        headers = ["Course Code", "Course Name", "Credit Hrs", "Max", "Passing", "Obtained", "Grade", "Grade Point"]

    for col_idx, h in enumerate(headers, start=1):
        c = ws.cell(row=row, column=col_idx, value=h)
        c.font = header_font
        c.fill = header_fill
        c.alignment = center
        c.border = thin_border
    row += 1

    for r in subj_rows:
        if inst_type == "School":
            values = [r["subject"], r["max"], r["passing"], r["obtained"], r["grade"], r["status"]]
        else:
            values = [
                r["course_code"], r["subject"], r["credit_hours"], r["max"],
                r["passing"], r["obtained"], r["grade"], r["gp"],
            ]
        for col_idx, v in enumerate(values, start=1):
            c = ws.cell(row=row, column=col_idx, value=v)
            c.border = thin_border
            c.alignment = center
        row += 1
    row += 1

    # ---- SUMMARY ----
    ws.cell(row=row, column=1, value="Summary").font = section_font
    row += 1

    if inst_type == "School":
        total_max = len(subj_rows) * 100
        summary_pairs = [
            ("Total Obtained Marks", st.session_state.total),
            ("Total Maximum Marks", total_max),
            ("Overall Percentage", f"{st.session_state.percentage:.2f}%"),
            ("Overall Grade", get_overall_grade(st.session_state.percentage)),
            ("Overall Result", overall_status),
        ]
    else:
        total_gp = sum(r["gp"] for r in subj_rows)
        gpa = total_gp / len(subj_rows)
        total_credit_hours = sum(r["credit_hours"] for r in subj_rows)
        summary_pairs = [
            ("Total Credit Hours", total_credit_hours),
            ("Total Grade Points", round(total_gp, 1)),
            ("Semester GPA", round(gpa, 2)),
            ("Overall Percentage", f"{st.session_state.percentage:.2f}%"),
            ("Overall Result", overall_status),
        ]

    for label, value in summary_pairs:
        ws.cell(row=row, column=1, value=label).font = label_font
        result_cell = ws.cell(row=row, column=2, value=value)
        if label == "Overall Result":
            result_cell.font = Font(bold=True, color="008000" if overall_status == "PASS" else "FF0000")
        row += 1
    row += 1

    # ---- REMARK (School) / ACADEMIC STANDING (University) ----
    if inst_type == "School":
        ws.cell(row=row, column=1, value="Teacher's Remark").font = section_font
        row += 1
        ws.cell(row=row, column=1, value=get_teacher_remark(st.session_state.percentage))
    else:
        total_gp = sum(r["gp"] for r in subj_rows)
        gpa = total_gp / len(subj_rows)
        ws.cell(row=row, column=1, value="Academic Standing").font = section_font
        row += 1
        ws.cell(row=row, column=1, value=get_academic_standing(gpa))

    # ---- Column widths for readability ----
    widths = [24, 22, 14, 12, 12, 12, 10, 12]
    for i in range(1, max_cols + 1):
        ws.column_dimensions[get_column_letter(i)].width = widths[i - 1]

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


# ==========================================================
# CALCULATE RESULT — report card, analytics, PDF/Excel export,
# print, and save-to-database. Only called once a result exists.
# ==========================================================
def render_result_report_section():

    st.divider()

    inst_type = st.session_state.institution_type
    subj_rows, all_pass = build_subject_rows(
        inst_type, st.session_state.subjects, st.session_state.marks
    )

    overall_status = "PASS" if all_pass else "FAIL"

    # Phase 3.1: pre-computed once here so the "Save to Database" section
    # can reuse them later without duplicating any calculation. The
    # School/University branches below still compute their own local
    # copies exactly as before (untouched) — these are simply the same
    # formulas evaluated once more for reuse outside those branches.
    overall_grade = get_overall_grade(st.session_state.percentage)
    gpa = None
    standing = None
    remark = None
    if inst_type == "University":
        _total_gp_precalc = sum(r["gp"] for r in subj_rows)
        gpa = _total_gp_precalc / len(subj_rows)
        standing = get_academic_standing(gpa)
    else:
        remark = get_teacher_remark(st.session_state.percentage)

    # ======================================================
    # SCHOOL — PROFESSIONAL BOARD-STYLE MARKSHEET (Phase 2.2)
    # UNCHANGED
    # ======================================================
    if inst_type == "School":

        overall_grade = get_overall_grade(st.session_state.percentage)
        remark = get_teacher_remark(st.session_state.percentage)
        total_max = len(subj_rows) * 100

        # ---------- STEP 1: Professional School Header ----------
        with st.container(border=True):
            st.markdown(f"### {st.session_state.institution_name or 'School Name'}")
            st.markdown("**Student Marksheet**")
            h1, h2, h3 = st.columns(3)
            h1.write(f"**Academic Session:** {st.session_state.academic_session}")
            h2.write(f"**Issue Date:** {st.session_state.issue_date}")
            h3.write(f"**Report No:** {st.session_state.report_number}")

        st.write("")

        # ---------- STEP 2: Student Information Box ----------
        with st.expander("Student Information", expanded=True):
            col1, col2 = st.columns(2)
            with col1:
                st.write("**Student Name:**", st.session_state.name)
                st.write("**Roll Number:**", st.session_state.roll_no)
                st.write("**Class:**", st.session_state.student_class)
            with col2:
                st.write("**Section:**", st.session_state.section)
                st.write("**Institution:**", st.session_state.institution_name)
                st.write("**Academic Session:**", st.session_state.academic_session)

        # ---------- STEP 3: Professional Subject Table ----------
        with st.expander("Subject-wise Marks Detail", expanded=True):
            table_display = [
                {
                    "Subject": r["subject"],
                    "Maximum Marks": r["max"],
                    "Passing Marks": r["passing"],
                    "Obtained Marks": r["obtained"],
                    "Result": r["status"]
                } for r in subj_rows
            ]
            st.table(table_display)

        # ---------- STEP 4: Marks Summary Box ----------
        st.subheader("Marks Summary")
        s1, s2, s3, s4 = st.columns(4)
        s1.metric("Total Obtained", f"{st.session_state.total}")
        s2.metric("Total Maximum", f"{total_max}")
        s3.metric("Percentage", f"{st.session_state.percentage:.2f}%")
        s4.metric("Overall Grade", overall_grade)

        if overall_status == "PASS":
            st.success(f"Overall Result: **{overall_status}**")
        else:
            st.error(f"Overall Result: **{overall_status}**")

        # ---------- STEP 5: Automatic Teacher Remarks ----------
        st.info(f"**Teacher's Remark:** {remark}")

        # ---------- STEP 6: Signature Section ----------
        with st.expander("Signatures", expanded=False):
            sig1, sig2, sig3 = st.columns(3)
            sig1.markdown("____________________")
            sig1.write("Teacher Signature")
            sig2.markdown("____________________")
            sig2.write("Principal Signature")
            sig3.markdown("____________________")
            sig3.write("School Stamp")

    # ======================================================
    # UNIVERSITY — PROFESSIONAL ACADEMIC TRANSCRIPT (Phase 2.3)
    # ======================================================
    else:

        total_gp = sum(r["gp"] for r in subj_rows)
        gpa = total_gp / len(subj_rows)
        total_credit_hours = sum(r["credit_hours"] for r in subj_rows)
        standing = get_academic_standing(gpa)

        # ---------- STEP 1: Professional University Header ----------
        with st.container(border=True):
            st.markdown(f"### {st.session_state.institution_name or 'University Name'}")
            st.markdown("**Official Academic Transcript**")
            h1, h2, h3 = st.columns(3)
            h1.write(f"**Transcript No:** {st.session_state.transcript_number}")
            h2.write(f"**Issue Date:** {st.session_state.uni_issue_date}")
            h3.write(f"**Academic Session:** {st.session_state.university_academic_session}")

        st.write("")

        # ---------- STEP 2: Student Information Section ----------
        with st.expander("Student Information", expanded=True):
            col1, col2 = st.columns(2)
            with col1:
                st.write("**Student Name:**", st.session_state.name)
                st.write("**Roll Number:**", st.session_state.roll_no)
                st.write("**Department:**", st.session_state.department)
            with col2:
                st.write("**Semester:**", st.session_state.semester)
                st.write("**Institution:**", st.session_state.institution_name)
                st.write("**Academic Session:**", st.session_state.university_academic_session)

        # ---------- STEP 3: Professional Transcript Table ----------
        with st.expander("Transcript — Course-wise Performance", expanded=True):
            table_display = [
                {
                    "Course Code": r["course_code"],
                    "Course Name": r["subject"],
                    "Credit Hrs": r["credit_hours"],
                    "Max Marks": r["max"],
                    "Passing Marks": r["passing"],
                    "Obtained Marks": r["obtained"],
                    "Letter Grade": r["grade"],
                    "Grade Point": r["gp"],
                } for r in subj_rows
            ]
            st.table(table_display)

        # ---------- STEP 4: Academic Performance Summary ----------
        st.subheader("Academic Performance Summary")
        c1, c2, c3, c4, c5 = st.columns(5)
        c1.metric("Credit Hours", total_credit_hours)
        c2.metric("Grade Points", f"{total_gp:.1f}")
        c3.metric("Semester GPA", f"{gpa:.2f}")
        c4.metric("Percentage", f"{st.session_state.percentage:.2f}%")
        c5.metric("Institution", st.session_state.institution_name or "—")

        if overall_status == "PASS":
            st.success(f"Overall Result: **{overall_status}**")
        else:
            st.error(f"Overall Result: **{overall_status}**")

        # ---------- STEP 5: Academic Standing ----------
        st.info(f"**Academic Standing:** {standing}")

        # ---------- STEP 6: Signature Section ----------
        with st.expander("Signatures", expanded=False):
            sig1, sig2, sig3 = st.columns(3)
            sig1.markdown("____________________")
            sig1.write("Controller of Examination")
            sig2.markdown("____________________")
            sig2.write("Registrar")
            sig3.markdown("____________________")
            sig3.write("University Seal")

    # ======================================================
    # ANALYTICS DASHBOARD (Phase 2.6) — appears below the report card
    # for both School and University. Display-only; does not touch
    # calculations, PDF generation, or session state used elsewhere.
    # Phase 2.7: gated by the sidebar "Show Analytics" / "Show Charts"
    # toggles — purely presentational, no logic changes.
    # ======================================================
    if st.session_state.show_analytics:
        render_analytics_dashboard(
            inst_type, subj_rows, overall_status,
            show_charts=st.session_state.show_charts,
        )
    else:
        st.info("Analytics dashboard is hidden (enable in sidebar Settings).")

    if st.session_state.dev_mode:
        with st.expander("Developer Mode — Internal State", expanded=False):
            st.caption("Read-only debug view. Does not affect calculations.")
            st.json({
                "institution_type": inst_type,
                "overall_status": overall_status,
                "total": st.session_state.total,
                "percentage": st.session_state.percentage,
                "subject_count": len(subj_rows),
            })

    # ----------------------------------------------------
    # PDF Generation — structured marksheet using reportlab tables
    # ----------------------------------------------------
    file_path = tempfile.NamedTemporaryFile(delete=False, suffix=".pdf").name
    doc = SimpleDocTemplate(
        file_path, pagesize=letter,
        topMargin=0.6 * inch, bottomMargin=0.95 * inch,
        leftMargin=0.7 * inch, rightMargin=0.7 * inch
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "TitleCenter", parent=styles["Heading1"],
        alignment=TA_CENTER, fontSize=18, spaceAfter=2,
        textColor=colors.HexColor("#1a1a2e")
    )
    subtitle_style = ParagraphStyle(
        "SubtitleCenter", parent=styles["Normal"],
        alignment=TA_CENTER, fontSize=11, textColor=colors.HexColor("#555555"),
        spaceAfter=10
    )
    section_style = ParagraphStyle(
        "Section", parent=styles["Heading2"], fontSize=13,
        textColor=colors.HexColor("#1a1a2e"), spaceBefore=14, spaceAfter=6
    )
    meta_style = ParagraphStyle(
        "MetaCenter", parent=styles["Normal"],
        alignment=TA_CENTER, fontSize=9, textColor=colors.HexColor("#777777"),
        spaceAfter=10
    )
    remark_style = ParagraphStyle(
        "Remark", parent=styles["Normal"], fontSize=11,
        textColor=colors.HexColor("#1a1a2e")
    )

    # University-only styles — kept separate from the School styles above
    # so the School PDF's look is never touched.
    university_title_style = ParagraphStyle(
        "UnivTitle", parent=styles["Heading1"],
        alignment=TA_CENTER, fontSize=20, spaceAfter=2,
        textColor=colors.HexColor("#12263a")
    )
    university_subtitle_style = ParagraphStyle(
        "UnivSubtitle", parent=styles["Normal"],
        alignment=TA_CENTER, fontSize=11, fontName="Helvetica-Bold",
        textColor=colors.HexColor("#b8860b"), spaceAfter=10
    )
    university_meta_style = ParagraphStyle(
        "UnivMeta", parent=styles["Normal"],
        alignment=TA_CENTER, fontSize=9, textColor=colors.HexColor("#666666"),
        spaceAfter=10
    )
    university_section_style = ParagraphStyle(
        "UnivSection", parent=styles["Heading2"], fontSize=13,
        textColor=colors.HexColor("#12263a"), spaceBefore=14, spaceAfter=6
    )
    standing_style = ParagraphStyle(
        "Standing", parent=styles["Normal"], fontSize=11,
        textColor=colors.HexColor("#12263a")
    )

    elements = []

    if inst_type == "School":
        # ---------- SCHOOL: Professional Board-Style Marksheet (UNCHANGED) ----------
        overall_grade = get_overall_grade(st.session_state.percentage)
        remark = get_teacher_remark(st.session_state.percentage)
        total_max = len(subj_rows) * 100

        # STEP 1 — Header (logo placeholder + title block)
        header_title_block = [
            Paragraph(st.session_state.institution_name or "School Name", title_style),
            Paragraph("Student Marksheet", subtitle_style),
            Paragraph(
                f"Academic Session: {st.session_state.academic_session} &nbsp;&nbsp;|&nbsp;&nbsp; "
                f"Issue Date: {st.session_state.issue_date} &nbsp;&nbsp;|&nbsp;&nbsp; "
                f"Report No: {st.session_state.report_number}",
                meta_style
            ),
        ]
        header_table = Table(
            [[draw_logo_placeholder("SCHOOL<br/>LOGO", "#1a1a2e"), header_title_block]],
            colWidths=[60, 400]
        )
        header_table.setStyle(TableStyle([
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("ALIGN", (0, 0), (0, 0), "CENTER"),
        ]))
        elements.append(header_table)
        elements.append(HRFlowable(width="100%", color=colors.HexColor("#1a1a2e"), thickness=1.2))
        elements.append(Spacer(1, 10))

        # STEP 2 — Student Information Box
        elements.append(Paragraph("Student Information", section_style))
        info_data = [
            ["Student Name:", st.session_state.name, "Roll No:", st.session_state.roll_no],
            ["Class:", st.session_state.student_class, "Section:", st.session_state.section],
            ["Institution:", st.session_state.institution_name, "Session:", st.session_state.academic_session],
            ["Issue Date:", st.session_state.issue_date, "Report No:", st.session_state.report_number],
        ]
        info_table = Table(info_data, colWidths=[90, 180, 70, 110])
        info_table.setStyle(TableStyle([
            ("FONTNAME", (0, 0), (-1, -1), "Helvetica"),
            ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"),
            ("FONTNAME", (2, 0), (2, -1), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 10),
            ("BOX", (0, 0), (-1, -1), 0.8, colors.HexColor("#999999")),
            ("INNERGRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#cccccc")),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
            ("TOPPADDING", (0, 0), (-1, -1), 7),
            ("LEFTPADDING", (0, 0), (-1, -1), 8),
        ]))
        elements.append(info_table)
        elements.append(Spacer(1, 14))

        # STEP 3 — Professional Subject Table
        elements.append(Paragraph("Subject-wise Marks", section_style))
        table_data = [["Subject", "Maximum Marks", "Passing Marks", "Obtained Marks", "Result"]]
        for r in subj_rows:
            table_data.append([
                r["subject"], str(r["max"]), str(r["passing"]),
                str(r["obtained"]), r["status"]
            ])

        marks_table = Table(table_data, colWidths=[150, 85, 85, 90, 70], repeatRows=1)
        marks_table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1a1a2e")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTNAME", (0, 1), (-1, -1), "Helvetica"),
            ("FONTSIZE", (0, 0), (-1, -1), 10),
            ("ALIGN", (1, 0), (-1, -1), "CENTER"),
            ("BOX", (0, 0), (-1, -1), 0.8, colors.HexColor("#1a1a2e")),
            ("GRID", (0, 0), (-1, -1), 0.6, colors.HexColor("#999999")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f2f2f7")]),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
            ("TOPPADDING", (0, 0), (-1, -1), 6),
        ]))
        elements.append(marks_table)
        elements.append(Spacer(1, 14))

        # STEP 4 — Marks Summary Box
        elements.append(Paragraph("Marks Summary", section_style))
        summary_data = [
            ["Total Obtained Marks:", f"{st.session_state.total}", "Total Maximum Marks:", f"{total_max}"],
            ["Overall Percentage:", f"{st.session_state.percentage:.2f}%", "Overall Grade:", overall_grade],
            ["Overall Result:", overall_status, "", ""],
        ]
        summary_table = Table(summary_data, colWidths=[120, 90, 130, 100])
        summary_table.setStyle(TableStyle([
            ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"),
            ("FONTNAME", (2, 0), (2, -1), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 10.5),
            ("BOX", (0, 0), (-1, -1), 0.8, colors.HexColor("#999999")),
            ("INNERGRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#cccccc")),
            ("TEXTCOLOR", (1, 2), (1, 2),
             colors.green if overall_status == "PASS" else colors.red),
            ("FONTNAME", (1, 2), (1, 2), "Helvetica-Bold"),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
            ("TOPPADDING", (0, 0), (-1, -1), 7),
            ("LEFTPADDING", (0, 0), (-1, -1), 8),
        ]))
        elements.append(summary_table)
        elements.append(Spacer(1, 14))

        # STEP 5 — Automatic Teacher Remarks
        elements.append(Paragraph("Teacher's Remark", section_style))
        remark_table = Table([[Paragraph(f"<b>{remark}</b>", remark_style)]], colWidths=[440])
        remark_table.setStyle(TableStyle([
            ("BOX", (0, 0), (-1, -1), 0.8, colors.HexColor("#1a1a2e")),
            ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#f2f2f7")),
            ("TOPPADDING", (0, 0), (-1, -1), 10),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 10),
            ("LEFTPADDING", (0, 0), (-1, -1), 10),
        ]))
        elements.append(remark_table)
        elements.append(Spacer(1, 36))

        # STEP 6 — Signature Section
        sig_data = [
            ["_______________________", "_______________________", "_______________________"],
            ["Teacher Signature", "Principal Signature", "School Stamp"],
        ]
        sig_table = Table(sig_data, colWidths=[150, 150, 150])
        sig_table.setStyle(TableStyle([
            ("ALIGN", (0, 0), (-1, -1), "CENTER"),
            ("FONTSIZE", (0, 0), (-1, -1), 10),
            ("FONTNAME", (0, 1), (-1, 1), "Helvetica-Bold"),
            ("TOPPADDING", (0, 1), (-1, 1), 4),
        ]))
        elements.append(sig_table)
        elements.append(Spacer(1, 20))
        elements.append(HRFlowable(width="100%", color=colors.HexColor("#cccccc"), thickness=0.7))
        elements.append(Spacer(1, 4))
        elements.append(Paragraph(
            "This is a computer-generated mark sheet.", styles["Normal"]
        ))

    else:
        # ---------- UNIVERSITY: Professional Academic Transcript (Phase 2.3) ----------
        total_gp = sum(r["gp"] for r in subj_rows)
        gpa = total_gp / len(subj_rows)
        total_credit_hours = sum(r["credit_hours"] for r in subj_rows)
        standing = get_academic_standing(gpa)

        # STEP 1 — Professional University Header (logo placeholder + title block)
        header_title_block = [
            Paragraph(st.session_state.institution_name or "University Name", university_title_style),
            Paragraph("OFFICIAL ACADEMIC TRANSCRIPT", university_subtitle_style),
            Paragraph(
                f"Transcript No: {st.session_state.transcript_number} &nbsp;&nbsp;|&nbsp;&nbsp; "
                f"Issue Date: {st.session_state.uni_issue_date} &nbsp;&nbsp;|&nbsp;&nbsp; "
                f"Academic Session: {st.session_state.university_academic_session}",
                university_meta_style
            ),
        ]
        header_table = Table(
            [[draw_logo_placeholder("UNIVERSITY<br/>LOGO", "#12263a"), header_title_block]],
            colWidths=[60, 400]
        )
        header_table.setStyle(TableStyle([
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("ALIGN", (0, 0), (0, 0), "CENTER"),
        ]))
        elements.append(header_table)
        elements.append(HRFlowable(width="100%", color=colors.HexColor("#b8860b"), thickness=1.5))
        elements.append(Spacer(1, 10))

        # STEP 2 — Student Information Section
        elements.append(Paragraph("Student Information", university_section_style))
        info_data = [
            ["Student Name:", st.session_state.name, "Roll No:", st.session_state.roll_no],
            ["Department:", st.session_state.department, "Semester:", st.session_state.semester],
            ["Institution:", st.session_state.institution_name, "Session:", st.session_state.university_academic_session],
            ["Issue Date:", st.session_state.uni_issue_date, "Transcript No:", st.session_state.transcript_number],
        ]
        info_table = Table(info_data, colWidths=[90, 180, 70, 110])
        info_table.setStyle(TableStyle([
            ("FONTNAME", (0, 0), (-1, -1), "Helvetica"),
            ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"),
            ("FONTNAME", (2, 0), (2, -1), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 10),
            ("BOX", (0, 0), (-1, -1), 0.8, colors.HexColor("#12263a")),
            ("INNERGRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#cccccc")),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
            ("TOPPADDING", (0, 0), (-1, -1), 7),
            ("LEFTPADDING", (0, 0), (-1, -1), 8),
        ]))
        elements.append(info_table)
        elements.append(Spacer(1, 14))

        # STEP 3 — Professional Transcript Table
        elements.append(Paragraph("Transcript — Course-wise Performance", university_section_style))
        table_data = [["Code", "Course Name", "Cr. Hrs", "Max", "Passing", "Obtained", "Grade", "GP"]]
        for r in subj_rows:
            table_data.append([
                r["course_code"], r["subject"], str(r["credit_hours"]),
                str(r["max"]), str(r["passing"]), str(r["obtained"]),
                r["grade"], f"{r['gp']:.1f}"
            ])

        marks_table = Table(table_data, colWidths=[55, 140, 45, 40, 55, 55, 45, 40], repeatRows=1)
        marks_table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#12263a")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#e8c874")),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTNAME", (0, 1), (-1, -1), "Helvetica"),
            ("FONTSIZE", (0, 0), (-1, -1), 9),
            ("ALIGN", (2, 0), (-1, -1), "CENTER"),
            ("BOX", (0, 0), (-1, -1), 0.8, colors.HexColor("#12263a")),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#999999")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f4f0e6")]),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
            ("TOPPADDING", (0, 0), (-1, -1), 6),
        ]))
        elements.append(marks_table)
        elements.append(Spacer(1, 14))

        # STEP 4 — Academic Performance Summary
        elements.append(Paragraph("Academic Performance Summary", university_section_style))
        summary_data = [
            ["Total Credit Hours:", str(total_credit_hours), "Total Grade Points:", f"{total_gp:.1f}"],
            ["Semester GPA:", f"{gpa:.2f}", "Overall Percentage:", f"{st.session_state.percentage:.2f}%"],
            ["Overall Result:", overall_status, "", ""],
        ]
        summary_table = Table(summary_data, colWidths=[120, 90, 130, 100])
        summary_table.setStyle(TableStyle([
            ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"),
            ("FONTNAME", (2, 0), (2, -1), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 10.5),
            ("BOX", (0, 0), (-1, -1), 0.8, colors.HexColor("#999999")),
            ("INNERGRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#cccccc")),
            ("TEXTCOLOR", (1, 2), (1, 2),
             colors.green if overall_status == "PASS" else colors.red),
            ("FONTNAME", (1, 2), (1, 2), "Helvetica-Bold"),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
            ("TOPPADDING", (0, 0), (-1, -1), 7),
            ("LEFTPADDING", (0, 0), (-1, -1), 8),
        ]))
        elements.append(summary_table)
        elements.append(Spacer(1, 14))

        # STEP 5 — Academic Standing
        elements.append(Paragraph("Academic Standing", university_section_style))
        standing_table = Table([[Paragraph(f"<b>{standing}</b>", standing_style)]], colWidths=[440])
        standing_table.setStyle(TableStyle([
            ("BOX", (0, 0), (-1, -1), 0.8, colors.HexColor("#b8860b")),
            ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#f4f0e6")),
            ("TOPPADDING", (0, 0), (-1, -1), 10),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 10),
            ("LEFTPADDING", (0, 0), (-1, -1), 10),
        ]))
        elements.append(standing_table)
        elements.append(Spacer(1, 36))

        # STEP 6 — Signature Section
        sig_data = [
            ["_______________________", "_______________________", "_______________________"],
            ["Controller of Examination", "Registrar", "University Seal"],
        ]
        sig_table = Table(sig_data, colWidths=[150, 150, 150])
        sig_table.setStyle(TableStyle([
            ("ALIGN", (0, 0), (-1, -1), "CENTER"),
            ("FONTSIZE", (0, 0), (-1, -1), 10),
            ("FONTNAME", (0, 1), (-1, 1), "Helvetica-Bold"),
            ("TOPPADDING", (0, 1), (-1, 1), 4),
        ]))
        elements.append(sig_table)
        elements.append(Spacer(1, 20))
        elements.append(HRFlowable(width="100%", color=colors.HexColor("#cccccc"), thickness=0.7))
        elements.append(Spacer(1, 4))
        elements.append(Paragraph(
            "This is an official computer-generated academic transcript.",
            styles["Normal"]
        ))

    if inst_type == "School":
        doc.build(elements, onFirstPage=add_school_footer, onLaterPages=add_school_footer)
    else:
        doc.build(elements, onFirstPage=add_university_footer, onLaterPages=add_university_footer)

    st.success("PDF generated successfully.")

    # ----------------------------------------------------
    # Phase 2.7 / STEP 1 — Excel Export
    # Built from the same subj_rows / session_state values as the PDF,
    # so both exports always stay in sync automatically.
    # ----------------------------------------------------
    excel_buffer = generate_excel_report(inst_type, subj_rows, overall_status)

    # ==========================================================
    # DOWNLOAD REPORT EXPANDER
    # PDF export (existing, unchanged) + Excel export (new) side by side.
    # ==========================================================
    with st.expander("Download Report", expanded=True):
        dl_col1, dl_col2 = st.columns(2)

        with dl_col1:
            with open(file_path, "rb") as file:
                st.download_button(
                    label=(
                        "Download Marksheet PDF" if inst_type == "School"
                        else "Download Academic Transcript PDF"
                    ),
                    data=file,
                    file_name=(
                        "Student_Marksheet.pdf" if inst_type == "School"
                        else "Academic_Transcript.pdf"
                    ),
                    mime="application/pdf",
                    use_container_width=True,
                )

        with dl_col2:
            if excel_buffer is not None:
                st.download_button(
                    label=(
                        "Download Marksheet Excel" if inst_type == "School"
                        else "Download Transcript Excel"
                    ),
                    data=excel_buffer,
                    file_name=(
                        "Student_Marksheet.xlsx" if inst_type == "School"
                        else "Academic_Transcript.xlsx"
                    ),
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    use_container_width=True,
                )
            else:
                st.warning(
                    "Excel export needs the 'openpyxl' package. "
                    "Install it with `pip install openpyxl`."
                )

    # ==========================================================
    # Phase 2.7 / STEP 2 — Print Support
    # Streamlit has no native "print to printer" API, so the
    # future-ready approach is to trigger the browser's own print
    # dialog (window.print()) from a small embedded HTML/JS button.
    # The user's browser then handles print/"Save as PDF" natively.
    # This never removes or replaces the Download PDF button above.
    # ==========================================================
    with st.expander("Print Report", expanded=False):
        st.caption(
            "Opens your browser's print dialog for this page. "
            "For a cleaner printout, use the PDF download above."
        )
        components.html(
            """
            <div style="text-align:center;">
                <button onclick="window.print()" style="
                    padding: 0.5rem 1.2rem;
                    border-radius: 8px;
                    border: 1px solid #999;
                    background-color: #1a1a2e;
                    color: white;
                    font-weight: 600;
                    cursor: pointer;
                ">Print This Report</button>
            </div>
            <!--
                Future enhancement placeholder:
                Once a dedicated print-friendly HTML template of the
                report is available, replace window.print() with logic
                that opens that template in a new tab/iframe and calls
                print() on it, so only the report (not the whole
                Streamlit UI) is sent to the printer.
            -->
            """,
            height=70,
        )

    # ==========================================================
    # Phase 3.1 / STEP 3 — Save Result to MongoDB
    # Uses the overall_grade/gpa/standing/remark pre-computed above and
    # the already-built subj_rows — no calculation happens here.
    # ==========================================================
    with st.expander("Save to Database", expanded=False):
        if not MONGODB_AVAILABLE:
            st.warning(
                "MongoDB integration isn't available in this environment — "
                "install `pymongo` (and optionally `python-dotenv`) to enable saving results."
            )
        else:
            db_connected, db_conn_msg = test_connection()
            if not db_connected:
                st.warning(f"Could not reach MongoDB: {db_conn_msg}")
                st.caption(
                    "Set the `MONGO_URI` environment variable to your MongoDB "
                    "connection string, then reload the app."
                )
            else:
                st.caption(
                    f"Database: `{config.MONGO_DB_NAME}` · "
                    f"Collections: `{config.STUDENTS_COLLECTION}`, `{config.RESULTS_COLLECTION}`"
                )

            if st.button("Save Result", use_container_width=True, disabled=not db_connected):
                save_success, save_message = save_student_result(
                    inst_type, st.session_state, subj_rows, overall_status,
                    overall_grade, remark=remark, gpa=gpa, academic_standing=standing,
                )
                if save_success:
                    st.success(f"{save_message}")
                else:
                    st.error(f"{save_message}")



# ==========================================================
# AI INSIGHTS — feedback, predictor & chatbot (GradeSense AI)
# ==========================================================
def render_ai_insights_section():
    # Recomputed independently so this page works standalone
    # (Admin may navigate here directly without re-running the
    # Calculate Result page first in the same rerun).
    inst_type = st.session_state.institution_type
    subj_rows, all_pass = build_subject_rows(
        inst_type, st.session_state.subjects, st.session_state.marks
    )
    overall_status = "PASS" if all_pass else "FAIL"
    overall_grade = get_overall_grade(st.session_state.percentage)
    gpa = None
    if inst_type == "University":
        _total_gp_precalc = sum(r["gp"] for r in subj_rows)
        gpa = _total_gp_precalc / len(subj_rows)

    # ======================================================
    # AI INSIGHTS (GradeSense AI) — feedback, trend predictor & chatbot.
    # Display-only: reads the already-calculated result/subj_rows and
    # never touches grading, PDF, or Excel logic.
    # ======================================================
    if AI_INSIGHTS_AVAILABLE:
        st.divider()
        st.subheader("AI Insights — GradeSense AI")

        tab_feedback, tab_predict, tab_chat = st.tabs(
            ["AI Feedback", "AI Predictor", "Ask GradeSense AI"]
        )

        # ---------- Tab 1: AI Feedback ----------
        with tab_feedback:
            with st.spinner("Generating personalized feedback..."):
                feedback_text = generate_ai_feedback(
                    name=st.session_state.get("name", "Student"),
                    subj_rows=subj_rows,
                    percentage=st.session_state.percentage,
                    overall_status=overall_status,
                    inst_type=inst_type,
                )
            st.markdown(feedback_text)

        # ---------- Tab 2: AI Predictor ----------
        with tab_predict:
            roll_no_for_predict = st.session_state.get("roll_no")
            get_roll_fn = get_student_by_roll if MONGODB_AVAILABLE else None
            prediction = predict_next_performance(
                roll_no_for_predict, MONGODB_AVAILABLE, get_roll_fn
            )
            if prediction["ok"]:
                trend_icon = {"improving": "", "declining": "", "steady": ""}[prediction["trend"]]
                c1, c2, c3 = st.columns(3)
                c1.metric("Predicted Next %", f"{prediction['predicted']}%")
                c2.metric("Trend", f"{trend_icon} {prediction['trend'].title()}")
                c3.metric("Past Results Used", len(prediction["history"]))
                st.line_chart(prediction["history"])
                st.caption(
                    "Prediction is based on a simple trend line over your previously "
                    "saved results for this Roll Number — treat it as a guide, not a guarantee."
                )
            else:
                st.info(prediction["message"])

        # ---------- Tab 3: Chatbot ----------
        with tab_chat:
            st.caption("Ask about your result — e.g. \"What's my weakest subject?\", \"Am I passing?\", \"How can I improve?\"")
            chat_context = build_chat_context(
                name=st.session_state.get("name", "Student"),
                roll_no=st.session_state.get("roll_no"),
                subj_rows=subj_rows,
                percentage=st.session_state.percentage,
                overall_status=overall_status,
                overall_grade=overall_grade,
                inst_type=inst_type,
                gpa=gpa,
            )

            if "ai_chat_history" not in st.session_state:
                st.session_state.ai_chat_history = []

            for role, msg in st.session_state.ai_chat_history:
                with st.chat_message(role):
                    st.markdown(msg)

            user_q = st.chat_input("Ask GradeSense AI about your result...")
            if user_q:
                st.session_state.ai_chat_history.append(("user", user_q))
                reply = answer_student_question(user_q, chat_context)
                st.session_state.ai_chat_history.append(("assistant", reply))
                st.rerun()



# ==========================================================
# STUDENT RECORDS — browse/search saved results (read-only)
# ==========================================================
def render_student_records_section():
    st.divider()
    st.subheader("Student Records")

    if not MONGODB_AVAILABLE:
        st.info(
            "MongoDB integration isn't available in this environment — "
            "install `pymongo` to enable viewing saved student records here."
        )
    else:
        records_connected, records_conn_msg = test_connection()
        if not records_connected:
            st.info(f"Not connected to MongoDB yet. ({records_conn_msg})")
        else:
            records_tab_all, records_tab_search = st.tabs(
                ["All Students", "Search by Roll Number"]
            )

            # ---------- View all students ----------
            with records_tab_all:
                if st.button("Refresh", key="refresh_students_btn"):
                    st.rerun()

                fetch_success, students_or_error = get_all_students()
                if not fetch_success:
                    st.error(f"Could not load records: {students_or_error}")
                elif not students_or_error:
                    st.info("No student records saved yet. Calculate a result and use "
                            "**Save to Database** above to add one.")
                else:
                    records_table = [
                        {
                            "Name": s.get("name"),
                            "Roll No": s.get("roll_no"),
                            "Institution": s.get("institution_name"),
                            "Type": s.get("institution_type"),
                            "Percentage": s.get("percentage"),
                            "Grade": s.get("overall_grade"),
                            "GPA": s.get("gpa"),
                            "Saved On": (
                                s["created_date"].strftime("%Y-%m-%d %H:%M")
                                if s.get("created_date") else ""
                            ),
                        }
                        for s in students_or_error
                    ]
                    st.dataframe(records_table, use_container_width=True, hide_index=True)

            # ---------- Search by roll number ----------
            with records_tab_search:
                search_roll_no = st.text_input("Enter Roll Number", key="search_roll_input")
                if st.button("Search", key="search_roll_button", use_container_width=True):
                    if not search_roll_no.strip():
                        st.warning("Please enter a roll number to search.")
                    else:
                        search_success, search_result = get_student_by_roll(search_roll_no.strip())
                        if not search_success:
                            st.error(f"Search failed: {search_result}")
                        elif not search_result:
                            st.info(f"No records found for roll number '{search_roll_no}'.")
                        else:
                            for record in search_result:
                                with st.container(border=True):
                                    st.markdown(f"**{record.get('name')}** — Roll No: {record.get('roll_no')}")
                                    rc1, rc2, rc3 = st.columns(3)
                                    rc1.metric("Percentage", f"{record.get('percentage', 0):.2f}%")
                                    rc2.metric("Grade", record.get("overall_grade", "—"))
                                    if record.get("institution_type") == "University":
                                        rc3.metric("GPA", f"{record.get('gpa', 0):.2f}" if record.get("gpa") else "—")
                                    else:
                                        rc3.metric("Remarks", record.get("remarks", "—"))

                                    st.caption(
                                        f"{record.get('institution_type')} · "
                                        f"{record.get('class_or_department')} · "
                                        f"{record.get('section_or_semester')} · "
                                        f"Session {record.get('academic_session')} · "
                                        f"Saved {record['created_date'].strftime('%Y-%m-%d %H:%M') if record.get('created_date') else ''}"
                                    )

                                    if record.get("subjects"):
                                        subject_detail_table = [
                                            {"Subject": subj, "Marks": mark, "Grade": grade}
                                            for subj, mark, grade in zip(
                                                record.get("subjects", []),
                                                record.get("marks", []),
                                                record.get("grades", []),
                                            )
                                        ]
                                        st.table(subject_detail_table)


# ==========================================================
# FOOTER
# ==========================================================
def render_footer():
    st.markdown(
        f"""
        <div class="sgc-footer">
            <b>{APP_NAME}</b><br>
            Developed using Python • Streamlit • ReportLab • Plotly • OpenPyXL<br>
            Version {APP_VERSION} &nbsp;|&nbsp; © {APP_YEAR}
        </div>
        """,
        unsafe_allow_html=True,
    )


# ---------- Phase 3.3 / STEP 5: Dashboard routing (RBAC) ----------
# Admin gets separate, dedicated pages (User Management, Student Info &
# Subjects, Calculate Result, AI Insights, Student Records) chosen via
# the sidebar navigation above. Teacher keeps the original single
# scrolling page (calculator + report + AI insights + records together).
# Student sees only their own dashboard.
user_role = st.session_state.get("role")

if user_role == "Admin":
    admin_page = st.session_state.get("admin_page", "User Management")

    if admin_page == "User Management":
        render_admin_dashboard()

    elif admin_page == "Student Info & Subjects":
        st.title("GradeSense AI")
        st.caption("Enter the student's information and subject-wise marks, then calculate the result.")
        render_data_entry_section()

    elif admin_page == "Calculate Result":
        st.title("GradeSense AI")
        st.caption("Marksheet / Transcript, analytics, and export options.")
        if st.session_state.result:
            render_result_report_section()
        else:
            st.info(
                "ℹNo result calculated yet. Go to **Student Info & Subjects**, "
                "fill in the details, and click **Calculate Grade** first."
            )

    elif admin_page == "AI Insights":
        st.title("GradeSense AI")
        st.caption("AI-generated feedback, a performance predictor, and a chatbot for the current result.")
        if st.session_state.result:
            render_ai_insights_section()
        else:
            st.info(
                "ℹCalculate a result first (in **Student Info & Subjects**) "
                "to see AI insights."
            )

    elif admin_page == "Student Records":
        st.title("GradeSense AI")
        render_student_records_section()

    render_footer()
    st.stop()

elif user_role == "Teacher":
    render_teacher_dashboard()
    render_header_about_help()
    render_data_entry_section()
    if st.session_state.result:
        render_result_report_section()
        render_ai_insights_section()
    render_student_records_section()
    render_footer()
    st.stop()

elif user_role == "Student":
    render_student_dashboard()
    st.stop()

else:
    st.error("Unknown role.")
    st.stop()
